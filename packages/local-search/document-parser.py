#!/usr/bin/env python3
"""Bounded, offline document text extraction. Input is the original file on stdin."""

from __future__ import annotations

import hashlib
import importlib.metadata
import io
import itertools
import json
import platform
import re
import resource
import stat
import sys
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import PurePosixPath
from typing import Any, BinaryIO, NoReturn
from xml.etree import ElementTree


VERSION = 1
SUPPORTED = {"pdf", "docx", "pptx", "xlsx"}
XML_FORBIDDEN = re.compile(r"<!\s*(?:DOCTYPE|ENTITY)\b", re.IGNORECASE)
XML_ENCODING = re.compile(r"^\s*<\?xml\s+[^>]*encoding\s*=\s*['\"]([^'\"]+)['\"]", re.IGNORECASE)
OFFICE_NAMESPACE = {
    "xlsx": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "word": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "drawing": "http://schemas.openxmlformats.org/drawingml/2006/main",
}
MAX_EXCEL_ROW = 1_048_576
MAX_EXCEL_COLUMN = 16_384
PDF_DECODED_STREAM_LIMIT = 4 * 1024 * 1024
PDF_DECODED_TOTAL_LIMIT = 16 * 1024 * 1024
MEMORY_LIMIT = 512 * 1024 * 1024
CPU_SECONDS_LIMIT = 20


class ParserError(Exception):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


def fail(code: str, message: str) -> NoReturn:
    raise ParserError(code, message)


def read_bounded(stream: BinaryIO, limit: int) -> bytes:
    data = stream.read(limit + 1)
    if len(data) > limit:
        fail("INPUT_LIMIT", "document exceeds the input limit")
    return data


def apply_resource_limits() -> bool:
    """Apply OS-enforced limits before importing document parser libraries."""
    memory_enforced = False
    for limit_name, value in (("RLIMIT_AS", MEMORY_LIMIT), ("RLIMIT_DATA", MEMORY_LIMIT)):
        if hasattr(resource, limit_name):
            limit = getattr(resource, limit_name)
            _, hard = resource.getrlimit(limit)
            bounded = value if hard < 0 or hard >= value else hard
            try:
                resource.setrlimit(limit, (bounded, hard))
                memory_enforced = True
            except (OSError, ValueError):
                continue
    soft, hard = resource.getrlimit(resource.RLIMIT_CPU)
    bounded_cpu = CPU_SECONDS_LIMIT if hard < 0 or hard >= CPU_SECONDS_LIMIT else hard
    resource.setrlimit(resource.RLIMIT_CPU, (bounded_cpu, hard))
    return memory_enforced


def decode_xml(data: bytes) -> str:
    if data.startswith((b"\x00\x00\xfe\xff", b"\xff\xfe\x00\x00")):
        fail("UNSUPPORTED_XML_ENCODING", "UTF-32 Office XML is unsupported")
    if data.startswith(b"\xef\xbb\xbf"):
        encoding = "utf-8-sig"
    elif data.startswith(b"\xff\xfe"):
        encoding = "utf-16-le"
        data = data[2:]
    elif data.startswith(b"\xfe\xff"):
        encoding = "utf-16-be"
        data = data[2:]
    elif data.startswith(b"\x3c\x00\x3f\x00"):
        encoding = "utf-16-le"
    elif data.startswith(b"\x00\x3c\x00\x3f"):
        encoding = "utf-16-be"
    else:
        encoding = "utf-8"
    try:
        text = data.decode(encoding, errors="strict")
    except UnicodeDecodeError as error:
        raise ParserError("UNSUPPORTED_XML_ENCODING", "Office XML encoding is invalid or unsupported") from error
    declaration = XML_ENCODING.match(text)
    if declaration:
        declared = declaration.group(1).casefold().replace("_", "-")
        allowed = {"utf-8", "utf8", "us-ascii", "ascii"} if encoding.startswith("utf-8") else {"utf-16", encoding}
        if declared not in allowed:
            fail("UNSUPPORTED_XML_ENCODING", "Office XML declares an unsupported encoding")
    if XML_FORBIDDEN.search(text):
        fail("UNSAFE_XML", "DTD and entity declarations are unsupported")
    return text


def xml_root(data: bytes) -> ElementTree.Element:
    try:
        return ElementTree.fromstring(decode_xml(data))
    except ParserError:
        raise
    except ElementTree.ParseError as error:
        raise ParserError("MALFORMED_ARCHIVE", "Office XML is malformed") from error


def cell_bounds(reference: str) -> tuple[int, int, int, int]:
    match = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?([1-9][0-9]*)(?::\$?([A-Za-z]{1,3})\$?([1-9][0-9]*))?", reference)
    if not match:
        fail("MALFORMED_ARCHIVE", "worksheet contains an invalid cell range")
    def column(value: str) -> int:
        result = 0
        for character in value.upper():
            result = result * 26 + ord(character) - 64
        return result
    first_column, first_row = column(match.group(1)), int(match.group(2))
    last_column = column(match.group(3) or match.group(1))
    last_row = int(match.group(4) or match.group(2))
    if first_column > last_column or first_row > last_row or last_column > MAX_EXCEL_COLUMN or last_row > MAX_EXCEL_ROW:
        fail("CELL_LIMIT", "worksheet range exceeds supported Excel bounds")
    return first_column, first_row, last_column, last_row


def range_area(reference: str) -> int:
    first_column, first_row, last_column, last_row = cell_bounds(reference)
    return (last_column - first_column + 1) * (last_row - first_row + 1)


def validate_office_zip(
    data: bytes,
    format_name: str,
    item_limit: int,
    cell_limit: int,
    member_limit: int,
    expanded_limit: int,
    member_bytes_limit: int,
    ratio_limit: int,
) -> list[str]:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except (OSError, zipfile.BadZipFile) as error:
        raise ParserError("MALFORMED_ARCHIVE", "invalid Office ZIP container") from error
    with archive:
        infos = archive.infolist()
        if len(infos) > member_limit:
            fail("ARCHIVE_LIMIT", "Office archive has too many members")
        names: set[str] = set()
        declared_total = 0
        actual_total = 0
        worksheet_cells = 0
        worksheet_extent = 0
        merged_cells = 0
        slide_count = 0
        limitations: set[str] = set()
        for info in infos:
            normalized = info.filename.replace("\\", "/")
            path = PurePosixPath(normalized)
            identity = normalized.casefold()
            if not normalized or normalized.startswith("/") or ".." in path.parts:
                fail("UNSAFE_ARCHIVE", "Office archive contains a traversal path")
            if identity in names:
                fail("UNSAFE_ARCHIVE", "Office archive contains duplicate member names")
            names.add(identity)
            if normalized.lower().endswith("vbaproject.bin") or "/macrosheets/" in f"/{normalized.lower()}/":
                fail("UNSUPPORTED_FORMAT", "macro-enabled Office documents are unsupported")
            if format_name == "pptx" and re.fullmatch(r"ppt/slides/slide\d+\.xml", normalized, re.IGNORECASE):
                slide_count += 1
                if slide_count > item_limit:
                    fail("SLIDE_LIMIT", "presentation exceeds the slide limit")
            if info.flag_bits & 0x1:
                fail("ENCRYPTED_UNSUPPORTED", "encrypted Office archives are unsupported")
            unix_mode = info.external_attr >> 16
            if unix_mode and stat.S_ISLNK(unix_mode):
                fail("UNSAFE_ARCHIVE", "Office archive contains a symbolic link")
            if info.file_size > member_bytes_limit:
                fail("ARCHIVE_LIMIT", "Office archive member exceeds its expanded limit")
            if info.file_size and (info.compress_size == 0 or info.file_size > info.compress_size * ratio_limit):
                fail("ARCHIVE_LIMIT", "Office archive compression ratio is unsafe")
            declared_total += info.file_size
            if declared_total > expanded_limit:
                fail("ARCHIVE_LIMIT", "Office archive exceeds its expanded limit")
            member = bytearray()
            try:
                with archive.open(info, "r") as source:
                    while True:
                        chunk = source.read(min(65_536, member_bytes_limit + 1 - len(member)))
                        if not chunk:
                            break
                        member.extend(chunk)
                        actual_total += len(chunk)
                        if len(member) > member_bytes_limit or actual_total > expanded_limit:
                            fail("ARCHIVE_LIMIT", "Office archive exceeded an actual decompression limit")
            except (OSError, RuntimeError, zipfile.BadZipFile) as error:
                raise ParserError("MALFORMED_ARCHIVE", "Office archive member cannot be read safely") from error
            if len(member) != info.file_size:
                fail("MALFORMED_ARCHIVE", "Office archive member size is inconsistent")
            root = None
            if normalized.lower().endswith((".xml", ".rels")):
                root = xml_root(bytes(member))
            if format_name == "xlsx" and re.fullmatch(r"xl/worksheets/sheet\d+\.xml", normalized, re.IGNORECASE):
                assert root is not None
                namespace = OFFICE_NAMESPACE["xlsx"]
                worksheet_cells += sum(1 for element in root.iter() if element.tag == f"{{{namespace}}}c")
                dimensions = [element.get("ref") for element in root.iter() if element.tag == f"{{{namespace}}}dimension"]
                for reference in dimensions:
                    if reference:
                        worksheet_extent += range_area(reference)
                for element in root.iter():
                    if element.tag == f"{{{namespace}}}mergeCell":
                        reference = element.get("ref")
                        if not reference:
                            fail("MALFORMED_ARCHIVE", "worksheet merge range is missing")
                        merged_cells += range_area(reference)
                if worksheet_cells > cell_limit or worksheet_extent > cell_limit or merged_cells > cell_limit:
                    fail("CELL_LIMIT", "workbook exceeds the cell limit")
            if format_name == "docx" and normalized.casefold() == "word/document.xml":
                assert root is not None
                word = OFFICE_NAMESPACE["word"]
                if any(ancestor.find(f".//{{{word}}}tbl") is not None for ancestor in root.findall(f".//{{{word}}}tbl")):
                    limitations.add("nestedTablesNotExtracted")
                if any(element.itertext() for element in root.findall(f".//{{{word}}}txbxContent")):
                    limitations.add("textBoxesNotExtracted")
            if format_name == "docx" and re.fullmatch(r"word/(?:header|footer)\d+\.xml", normalized, re.IGNORECASE):
                assert root is not None
                if any(text.strip() for text in root.itertext()):
                    limitations.add("headersFootersNotExtracted")
            if format_name == "docx" and normalized.casefold() == "word/footnotes.xml":
                assert root is not None
                if any(text.strip() for text in root.itertext()):
                    limitations.add("footnotesNotExtracted")
            if format_name == "docx" and normalized.casefold() == "word/endnotes.xml":
                assert root is not None
                if any(text.strip() for text in root.itertext()):
                    limitations.add("endnotesNotExtracted")
            if format_name == "docx" and normalized.casefold().startswith("word/media/"):
                limitations.add("imageTextNotExtracted")
            if format_name == "pptx" and re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", normalized, re.IGNORECASE):
                assert root is not None
                if any(text.strip() for text in root.itertext()):
                    limitations.add("notesNotExtracted")
            if format_name == "pptx" and re.fullmatch(r"ppt/charts/chart\d+\.xml", normalized, re.IGNORECASE):
                assert root is not None
                if any(text.strip() for text in root.itertext()):
                    limitations.add("chartTextNotExtracted")
            if format_name == "pptx" and normalized.casefold().startswith("ppt/media/"):
                limitations.add("imageTextNotExtracted")
        return sorted(limitations)


def utf16_length(value: str) -> int:
    return len(value.encode("utf-16-le")) // 2


@dataclass
class TextBuilder:
    utf8_limit: int
    segment_limit: int
    parts: list[str]
    segments: list[dict[str, Any]]
    utf8_bytes: int = 0
    utf16_units: int = 0

    @classmethod
    def create(cls, utf8_limit: int, segment_limit: int) -> TextBuilder:
        return cls(utf8_limit, segment_limit, [], [])

    def add(self, value: str | None, locator: dict[str, Any]) -> None:
        text = value or ""
        if not text.strip():
            return
        separator = "\n" if self.parts else ""
        added_bytes = len((separator + text).encode("utf-8"))
        if self.utf8_bytes + added_bytes > self.utf8_limit:
            fail("EXTRACTED_TEXT_LIMIT", "extracted UTF-8 text exceeds its limit")
        if len(self.segments) >= self.segment_limit:
            fail("SEGMENT_LIMIT", "document has too many text segments")
        if separator:
            self.parts.append(separator)
            self.utf8_bytes += 1
            self.utf16_units += 1
        start = self.utf16_units
        self.parts.append(text)
        self.utf8_bytes += len(text.encode("utf-8"))
        self.utf16_units += utf16_length(text)
        self.segments.append({"start": start, "end": self.utf16_units, "locator": locator})

    def text(self) -> str:
        return "".join(self.parts)


def preflight_pdf(reader: Any) -> list[str]:
    import pypdf.filters
    from pypdf.errors import LimitReachedError
    from pypdf.generic import IndirectObject, StreamObject

    pypdf.filters.ZLIB_MAX_OUTPUT_LENGTH = PDF_DECODED_STREAM_LIMIT
    supported_filters = {"/FlateDecode", "/Fl", "/ASCIIHexDecode", "/AHx", "/ASCII85Decode", "/A85", "/RunLengthDecode", "/RL"}
    limitations: set[str] = set()
    decoded_total = 0
    seen: set[tuple[int, int]] = set()
    xref = {generation: set(object_ids) for generation, object_ids in reader.xref.items()}
    xref.setdefault(0, set()).update(reader.xref_objStm)
    for generation, object_ids in xref.items():
        for object_id in object_ids:
            identity = (generation, object_id)
            if identity in seen:
                continue
            seen.add(identity)
            try:
                value = reader.get_object(IndirectObject(object_id, generation, reader))
            except LimitReachedError as error:
                raise ParserError("PDF_STREAM_LIMIT", "PDF decoded stream exceeds its safety limit") from error
            if not isinstance(value, StreamObject):
                continue
            subtype = str(value.get("/Subtype", ""))
            object_type = str(value.get("/Type", ""))
            if subtype == "/Image":
                limitations.add("imageTextNotExtracted")
                continue
            if subtype == "/Form":
                limitations.add("formXObjectCoverageUnverified")
            if object_type == "/XObject" and subtype not in {"", "/Form"}:
                limitations.add("unknownXObjectContent")
            filters = value.get("/Filter", [])
            if not isinstance(filters, list):
                filters = [filters]
            if any(str(filter_name) not in supported_filters for filter_name in filters):
                fail("UNSUPPORTED_PDF_CONTENT", "PDF uses a stream encoding outside the safety boundary")
            try:
                decoded = value.get_data()
            except LimitReachedError as error:
                raise ParserError("PDF_STREAM_LIMIT", "PDF decoded stream exceeds its safety limit") from error
            if len(decoded) > PDF_DECODED_STREAM_LIMIT:
                fail("PDF_STREAM_LIMIT", "PDF decoded stream exceeds its safety limit")
            if re.search(br"(?:^|\s)BI(?:\s)", decoded):
                limitations.add("imageTextNotExtracted")
            decoded_total += len(decoded)
            if decoded_total > PDF_DECODED_TOTAL_LIMIT:
                fail("PDF_STREAM_LIMIT", "PDF decoded streams exceed the aggregate safety limit")
    return sorted(limitations)


def extract_pdf(data: bytes, builder: TextBuilder, item_limit: int) -> tuple[bool, list[str]]:
    import pypdf.filters
    from pypdf import PdfReader
    try:
        pypdf.filters.ZLIB_MAX_OUTPUT_LENGTH = PDF_DECODED_STREAM_LIMIT
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted:
            fail("ENCRYPTED_UNSUPPORTED", "encrypted PDF files are unsupported")
        if len(reader.pages) > item_limit:
            fail("PAGE_LIMIT", "PDF exceeds the page limit")
        omitted = preflight_pdf(reader)
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            builder.add(text, {"page": number})
            contents = page.get_contents()
            if not text.strip() and contents is not None and contents.get_data().strip():
                omitted.append(f"page:{number}:nonTextContentNotExtracted")
        omitted = sorted(set(omitted))
        if omitted and not builder.segments:
            if "imageTextNotExtracted" in omitted:
                fail("OCR_REQUIRED", "PDF contains images but no extractable text")
            fail("UNSUPPORTED_PDF_CONTENT", "PDF contains non-text content outside the extraction coverage boundary")
        return not omitted, omitted
    except ParserError:
        raise
    except Exception as error:
        raise ParserError("MALFORMED_DOCUMENT", "PDF parsing failed") from error


def extract_docx(data: bytes, builder: TextBuilder) -> None:
    from docx import Document
    from docx.table import Table
    try:
        document = Document(io.BytesIO(data))
        paragraph_number = 0
        table_number = 0
        for block in document.iter_inner_content():
            if isinstance(block, Table):
                table_number += 1
                seen_cells: set[int] = set()
                for row_number, row in enumerate(block.rows, 1):
                    for cell_number, cell in enumerate(row.cells, 1):
                        identity = id(cell._tc)
                        if identity in seen_cells:
                            continue
                        seen_cells.add(identity)
                        for cell_paragraph, paragraph in enumerate(cell.paragraphs, 1):
                            builder.add(paragraph.text, {"table": table_number, "row": row_number, "cell": cell_number, "paragraph": cell_paragraph})
            else:
                paragraph_number += 1
                builder.add(block.text, {"paragraph": paragraph_number})
    except ParserError:
        raise
    except Exception as error:
        raise ParserError("MALFORMED_DOCUMENT", "DOCX parsing failed") from error


def walk_ppt_shapes(shapes: Any, slide_number: int, prefix: list[int], builder: TextBuilder) -> None:
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for shape_number, shape in enumerate(shapes, 1):
        shape_path = [*prefix, shape_number]
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            walk_ppt_shapes(shape.shapes, slide_number, shape_path, builder)
        elif getattr(shape, "has_table", False):
            seen_cells: set[int] = set()
            for row_number, row in enumerate(shape.table.rows, 1):
                for cell_number, cell in enumerate(row.cells, 1):
                    identity = id(cell._tc)
                    if identity in seen_cells:
                        continue
                    seen_cells.add(identity)
                    for paragraph_number, paragraph in enumerate(cell.text_frame.paragraphs, 1):
                        builder.add(paragraph.text, {"slide": slide_number, "shape": shape_path, "tableCell": {"row": row_number, "cell": cell_number}, "paragraph": paragraph_number})
        elif getattr(shape, "has_text_frame", False):
            for paragraph_number, paragraph in enumerate(shape.text_frame.paragraphs, 1):
                builder.add(paragraph.text, {"slide": slide_number, "shape": shape_path, "paragraph": paragraph_number})


def extract_pptx(data: bytes, builder: TextBuilder, item_limit: int) -> None:
    from pptx import Presentation
    try:
        presentation = Presentation(io.BytesIO(data))
        if len(presentation.slides) > item_limit:
            fail("SLIDE_LIMIT", "presentation exceeds the slide limit")
        for slide_number, slide in enumerate(presentation.slides, 1):
            walk_ppt_shapes(slide.shapes, slide_number, [], builder)
    except ParserError:
        raise
    except Exception as error:
        raise ParserError("MALFORMED_DOCUMENT", "PPTX parsing failed") from error


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def extract_xlsx(data: bytes, builder: TextBuilder, cell_limit: int) -> list[str]:
    from openpyxl import load_workbook
    try:
        formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        limitations: set[str] = set()
        cells_seen = 0
        try:
            for sheet_number, sheet in enumerate(formulas.worksheets):
                cached_sheet = cached.worksheets[sheet_number]
                for formula_row, cached_row in itertools.zip_longest(sheet.iter_rows(), cached_sheet.iter_rows(), fillvalue=()):
                    for cell, cached_cell in itertools.zip_longest(formula_row, cached_row, fillvalue=None):
                        if cell is None or cell.value is None:
                            continue
                        cells_seen += 1
                        if cells_seen > cell_limit:
                            fail("CELL_LIMIT", "workbook exceeds the cell limit")
                        locator: dict[str, Any] = {"sheet": sheet.title, "cell": cell.coordinate}
                        value = cell.value
                        if cell.data_type == "f":
                            limitations.add("cachedFormulaValuesMayBeStale")
                            value = cached_cell.value if cached_cell is not None else None
                            locator["valueSource"] = "cachedFormula"
                            if value is None:
                                limitations.add("formulaCacheMissing")
                        builder.add(cell_text(value), locator)
        finally:
            formulas.close()
            cached.close()
        return sorted(limitations)
    except ParserError:
        raise
    except Exception as error:
        raise ParserError("MALFORMED_DOCUMENT", "XLSX parsing failed") from error


def positive_limit(value: str, name: str) -> int:
    try:
        parsed = int(value)
    except ValueError as error:
        raise ParserError("INVALID_INVOCATION", f"invalid {name} limit") from error
    if parsed <= 0:
        fail("INVALID_INVOCATION", f"invalid {name} limit")
    return parsed


def parser_identity(format_name: str, limits: dict[str, int], memory_limit_enforced: bool) -> tuple[dict[str, Any], str]:
    distributions = {
        "pdf": "pypdf",
        "docx": "python-docx",
        "pptx": "python-pptx",
        "xlsx": "openpyxl",
    }
    identity = {
        "schemaVersion": VERSION,
        "implementation": "myos-local-document-parser",
        "implementationHash": hashlib.sha256(open(__file__, "rb").read()).hexdigest(),
        "python": platform.python_version(),
        "format": format_name,
        "libraries": {format_name: importlib.metadata.version(distributions[format_name])},
        "options": {
            "isolatedMode": True,
            "formulaEvaluation": False,
            "xmlEntities": False,
            "pdfDecodedStreamBytes": PDF_DECODED_STREAM_LIMIT,
            "pdfDecodedTotalBytes": PDF_DECODED_TOTAL_LIMIT,
            "addressSpaceBytes": MEMORY_LIMIT,
            "memoryLimitEnforced": memory_limit_enforced,
            "cpuSeconds": CPU_SECONDS_LIMIT,
        },
        "limits": limits,
    }
    serialized = json.dumps(identity, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return identity, hashlib.sha256(serialized).hexdigest()


def main() -> None:
    if len(sys.argv) != 11:
        fail("INVALID_INVOCATION", "expected format and nine limits")
    format_name = sys.argv[1].lower().lstrip(".")
    if format_name not in SUPPORTED:
        fail("UNSUPPORTED_FORMAT", "unsupported document format")
    input_limit = positive_limit(sys.argv[2], "input")
    text_limit = positive_limit(sys.argv[3], "text")
    item_limit = positive_limit(sys.argv[4], "page or slide")
    cell_limit = positive_limit(sys.argv[5], "cell")
    segment_limit = positive_limit(sys.argv[6], "segment")
    archive_member_limit = positive_limit(sys.argv[7], "archive member count")
    archive_expanded_limit = positive_limit(sys.argv[8], "archive expanded bytes")
    archive_member_bytes_limit = positive_limit(sys.argv[9], "archive member bytes")
    archive_ratio_limit = positive_limit(sys.argv[10], "archive compression ratio")
    memory_limit_enforced = apply_resource_limits()
    if format_name == "pdf" and not memory_limit_enforced:
        fail("UNSUPPORTED_SAFETY_BOUNDARY", "PDF parsing is disabled because this platform cannot enforce a child memory limit")
    data = read_bounded(sys.stdin.buffer, input_limit)
    builder = TextBuilder.create(text_limit, segment_limit)
    complete = True
    limitations: list[str] = []
    if format_name == "pdf":
        complete, limitations = extract_pdf(data, builder, item_limit)
    else:
        limitations = validate_office_zip(data, format_name, item_limit, cell_limit, archive_member_limit, archive_expanded_limit, archive_member_bytes_limit, archive_ratio_limit)
        if format_name == "docx":
            extract_docx(data, builder)
        elif format_name == "pptx":
            extract_pptx(data, builder, item_limit)
        elif format_name == "xlsx":
            limitations = sorted(set(limitations) | set(extract_xlsx(data, builder, cell_limit)))
        complete = not limitations
    coverage_status = "extracted" if complete else "partial"
    limits = {
        "inputBytes": input_limit,
        "extractedUtf8Bytes": text_limit,
        "pagesOrSlides": item_limit,
        "cells": cell_limit,
        "segments": segment_limit,
        "archiveMembers": archive_member_limit,
        "archiveExpandedBytes": archive_expanded_limit,
        "archiveMemberBytes": archive_member_bytes_limit,
        "archiveCompressionRatio": archive_ratio_limit,
    }
    identity, fingerprint = parser_identity(format_name, limits, memory_limit_enforced)
    result = {
        "version": VERSION,
        "sourceHash": hashlib.sha256(data).hexdigest(),
        "parserFingerprint": fingerprint,
        "parserIdentity": identity,
        "text": builder.text(),
        "segments": builder.segments,
        "complete": complete,
        "coverage": {"status": coverage_status, "limitations": limitations},
    }
    sys.stdout.write(json.dumps(result, ensure_ascii=False, separators=(",", ":")))


if __name__ == "__main__":
    try:
        main()
    except ParserError as error:
        sys.stderr.write(json.dumps({"error": {"code": error.code, "message": str(error)}}, separators=(",", ":")))
        raise SystemExit(2)
    except Exception:
        sys.stderr.write(json.dumps({"error": {"code": "PARSER_FAILED", "message": "document parser failed"}}, separators=(",", ":")))
        raise SystemExit(3)
